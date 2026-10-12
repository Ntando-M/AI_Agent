from pathlib import Path

from langchain_core.documents import Document

from docx import Document as DocxDocument
from pypdf import PdfReader

# TXT LOADER

def load_text_file(file_path : str) -> list[Document] :
    """
    Load a text file and return a Langchain Document.
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"File not found: {file_path}"
            )

    if path.suffix.lower() != ".txt":
        raise ValueError(
            f"Expected a TXT file, received: {path.suffix}"
            )
    
    text = path.read_text(
        encoding="utf-8"
    )

    return[
        Document(
            page_content=text,
            metadata={
                "source": str(path),
                "filename": path.name
                }
        )
    ]

# PDF LOADER

def load_pdf_file(file_path : str) -> list[Document] :
    """
    Load a PDF file and return a list of Langchain Documents, one per page.
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"File not found: {file_path}"
            )

    if path.suffix.lower() != ".pdf":
        raise ValueError(
            f"Expected a PDF file, received: {path.suffix}"
            )
    
    reader = PdfReader(
        str(path)
        )

    documents = []

    for page_number, page in enumerate(
        reader.pages,
        start=1
    ):

        text = page.extract_text() or ""

        if text:
            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": str(path),
                        "filename": path.name,
                        "page": page_number
                    }
                )
            )

    return documents

# DOCX LOADER

def load_docx_file(file_path : str) -> list[Document] :
    """
    Load a Word document and return a list of Langchain Documents,
    one per non-empty paragraph.

    DOCX has no page concept, so the paragraph index is recorded
    as "page" to keep the metadata shape identical to the PDF
    loader. Callers chunking on metadata do not need to know
    which format produced a Document.
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"File not found: {file_path}"
            )

    if path.suffix.lower() != ".docx":
        raise ValueError(
            f"Expected a DOCX file, received: {path.suffix}"
            )

    document = DocxDocument(
        str(path)
        )

    documents = []

    for index, paragraph in enumerate(
        document.paragraphs,
        start=1
        ):

        text = paragraph.text.strip()

        if text:
            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": str(path),
                        "filename": path.name,
                        "page": index
                    }
                )
            )

    return documents

# Document Directory loader

def load_documents(directory: str) -> list[Document] :
    """
    Load all documents from a directory and return a list of Langchain Documents.
    Supports TXT, PDF and DOCX files.

    The scan is recursive: the corpus is organised into
    per-format subdirectories (documents/pdf, documents/word),
    so a non-recursive glob finds nothing at the top level.
    """
    directory_path = Path(directory)

    if not directory_path.exists():
        raise FileNotFoundError(
            f"Directory not found: {directory}"
            )

    if not directory_path.is_dir():
        raise ValueError(
            f"Expected a directory, received: {directory_path}"
            )
    
    documents = []

# TXT Files

    for file_path in directory_path.rglob("*.txt"):

        documents.extend(
            load_text_file(
                str(file_path)
            )
        )

# PDF Files

    for file_path in directory_path.rglob("*.pdf"):

        documents.extend(
            load_pdf_file(
                str(file_path)
            )
        )

# DOCX Files

    for file_path in directory_path.rglob("*.docx"):

        documents.extend(
            load_docx_file(
                str(file_path)
            )
        )

    return documents
